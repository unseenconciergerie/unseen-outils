"""
Gestion du tableau Excel des prospects (assistant-prospection).

Le fichier est enregistré en dehors du dossier du projet, dans le dossier
iCloud "UNSEEN - Prospection", pour ne jamais dépendre d'un seul Mac.
"""

import os
from datetime import date

from openpyxl import Workbook, load_workbook

CHEMIN_DOSSIER_ICLOUD = "/Users/jaypandea/Documents/UNSEEN - Prospection"
CHEMIN_EXCEL = os.path.join(CHEMIN_DOSSIER_ICLOUD, "prospects.xlsx")

COLONNES = ["Pseudo", "Plateforme", "Prénom", "Email", "Date d'ajout", "Statut"]


def ouvrir_ou_creer_classeur():
    """Ouvre le tableau Excel, ou le crée avec ses colonnes s'il n'existe pas encore."""
    if os.path.exists(CHEMIN_EXCEL):
        return load_workbook(CHEMIN_EXCEL)

    os.makedirs(CHEMIN_DOSSIER_ICLOUD, exist_ok=True)
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = "Prospects"
    feuille.append(COLONNES)
    classeur.save(CHEMIN_EXCEL)
    return classeur


def trouver_doublon(pseudo, email):
    """Renvoie True si ce pseudo ou cet email est déjà dans le tableau."""
    classeur = ouvrir_ou_creer_classeur()
    feuille = classeur.active

    pseudo_recherche = pseudo.strip().lower()
    email_recherche = email.strip().lower()

    for ligne in feuille.iter_rows(min_row=2, values_only=True):
        pseudo_existant, _, _, email_existant, _, _ = ligne
        if str(pseudo_existant).strip().lower() == pseudo_recherche:
            return True
        if str(email_existant).strip().lower() == email_recherche:
            return True

    return False


def ajouter_ligne(pseudo, plateforme, prenom, email, statut):
    """Ajoute une ligne au tableau et enregistre le fichier."""
    classeur = ouvrir_ou_creer_classeur()
    feuille = classeur.active

    aujourdhui = date.today().strftime("%d/%m/%Y")
    feuille.append([pseudo, plateforme, prenom, email, aujourdhui, statut])
    classeur.save(CHEMIN_EXCEL)
